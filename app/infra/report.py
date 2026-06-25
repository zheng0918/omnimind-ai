"""导出报告渲染：结构化文档块 → docx / pdf 字节。

Python 侧不写 MinIO（只读源文件，见 clients.minio_client），故仅产出文件字节，
由 Java 解码 base64 后落桶 + 预签名下载。
- docx：python-docx；
- pdf：reportlab，中文走内置 CID 字体 STSong-Light（无需随包字体文件）。

块模型 Block = (kind, text)，kind ∈ {"h1", "h2", "p"}。
"""

from __future__ import annotations

import io

from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

DOCX_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)
PDF_CONTENT_TYPE = "application/pdf"

Block = tuple[str, str]

_PDF_FONT = "STSong-Light"
_pdf_font_ready = False


def markdown_to_blocks(md: str) -> list[Block]:
    """把导出用的简单 markdown 拆成块：# → h1，##/### → h2，其余非空行 → p。"""
    blocks: list[Block] = []
    for raw in (md or "").splitlines():
        line = raw.strip()
        if not line:
            continue
        if line.startswith("# "):
            blocks.append(("h1", line[2:].strip()))
        elif line.startswith("## "):
            blocks.append(("h2", line[3:].strip()))
        elif line.startswith("### "):
            blocks.append(("h2", line[4:].strip()))
        else:
            blocks.append(("p", line))
    return blocks


def render(title: str, blocks: list[Block], fmt: str) -> tuple[bytes, str, str]:
    """渲染为 (字节, content_type, 扩展名)。fmt ∈ {docx, pdf}，未知回退 docx。"""
    if (fmt or "docx").lower() == "pdf":
        return _render_pdf(title, blocks), PDF_CONTENT_TYPE, "pdf"
    return _render_docx(title, blocks), DOCX_CONTENT_TYPE, "docx"


def _render_docx(title: str, blocks: list[Block]) -> bytes:
    doc = Document()
    doc.add_heading(title, level=0)
    for kind, text in blocks:
        if kind == "h1":
            doc.add_heading(text, level=1)
        elif kind == "h2":
            doc.add_heading(text, level=2)
        else:
            doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _ensure_pdf_font() -> None:
    """进程内只注册一次 CJK 字体。"""
    global _pdf_font_ready
    if not _pdf_font_ready:
        pdfmetrics.registerFont(UnicodeCIDFont(_PDF_FONT))
        _pdf_font_ready = True


def _render_pdf(title: str, blocks: list[Block]) -> bytes:
    _ensure_pdf_font()
    styles = {
        "title": ParagraphStyle(
            "title", fontName=_PDF_FONT, fontSize=20, leading=28, spaceAfter=12
        ),
        "h1": ParagraphStyle(
            "h1", fontName=_PDF_FONT, fontSize=16, leading=22, spaceBefore=10, spaceAfter=6
        ),
        "h2": ParagraphStyle(
            "h2", fontName=_PDF_FONT, fontSize=13, leading=18, spaceBefore=8, spaceAfter=4
        ),
        "p": ParagraphStyle(
            "p", fontName=_PDF_FONT, fontSize=10.5, leading=16, spaceAfter=4
        ),
    }
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        topMargin=20 * mm,
        bottomMargin=20 * mm,
        leftMargin=20 * mm,
        rightMargin=20 * mm,
    )
    story: list[object] = [Paragraph(_esc(title), styles["title"]), Spacer(1, 6)]
    for kind, text in blocks:
        story.append(Paragraph(_esc(text), styles.get(kind, styles["p"])))
    doc.build(story)
    return buf.getvalue()


def _esc(text: str) -> str:
    """reportlab Paragraph 使用类 XML 标记，需转义 & < >。"""
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
