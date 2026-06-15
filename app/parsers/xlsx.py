"""XLSX 解析器（openpyxl）。每行非空单元格拼为一个段落。"""

from __future__ import annotations

import asyncio
import io

from openpyxl import load_workbook

from app.core.errors import CODE_PARSE_FAILED, ParseError
from app.schemas.parse import ParsedDocument, ParsedParagraph
from app.utils.text import normalize_text


class XlsxParser:
    """Excel 解析。逐 sheet 逐行；paragraph_id 记录 sheet 与行号。"""

    @property
    def supported_mime_types(self) -> tuple[str, ...]:
        return (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    async def parse(self, data: bytes) -> ParsedDocument:
        return await asyncio.to_thread(self._parse_sync, data)

    def _parse_sync(self, data: bytes) -> ParsedDocument:
        try:
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as exc:
            raise ParseError(CODE_PARSE_FAILED, "XLSX 解析失败：文件损坏或格式不支持") from exc
        paragraphs: list[ParsedParagraph] = []
        offset = 0
        try:
            for sheet in workbook.worksheets:
                for row_idx, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                    cells = [str(c) for c in row if c is not None and str(c).strip()]
                    if not cells:
                        continue
                    text = normalize_text(" | ".join(cells)).strip()
                    if not text:
                        continue
                    paragraphs.append(
                        ParsedParagraph(
                            text=text,
                            page=None,
                            paragraph_id=f"{sheet.title}!{row_idx}",
                            char_offset=offset,
                        )
                    )
                    offset += len(text)
        finally:
            workbook.close()
        if not paragraphs:
            raise ParseError(CODE_PARSE_FAILED, "XLSX 无可提取文本")
        return ParsedDocument(paragraphs=paragraphs, page_count=0)
