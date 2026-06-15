"""DOCX 解析器（python-docx）。按段落提取。"""

from __future__ import annotations

import asyncio
import io

from docx import Document

from app.core.errors import CODE_PARSE_FAILED, ParseError
from app.schemas.parse import ParsedDocument, ParsedParagraph
from app.utils.text import normalize_text


class DocxParser:
    """Word 文档解析。docx 无页码概念，page 置空。"""

    @property
    def supported_mime_types(self) -> tuple[str, ...]:
        return (
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )

    async def parse(self, data: bytes) -> ParsedDocument:
        return await asyncio.to_thread(self._parse_sync, data)

    def _parse_sync(self, data: bytes) -> ParsedDocument:
        try:
            document = Document(io.BytesIO(data))
        except Exception as exc:
            raise ParseError(CODE_PARSE_FAILED, "DOCX 解析失败：文件损坏或格式不支持") from exc
        paragraphs: list[ParsedParagraph] = []
        offset = 0
        seq = 0
        for para in document.paragraphs:
            text = normalize_text(para.text).strip()
            if not text:
                continue
            paragraphs.append(
                ParsedParagraph(
                    text=text, page=None, paragraph_id=f"p{seq}", char_offset=offset
                )
            )
            offset += len(text)
            seq += 1
        if not paragraphs:
            raise ParseError(CODE_PARSE_FAILED, "DOCX 无可提取文本")
        return ParsedDocument(paragraphs=paragraphs, page_count=0)
